"""Compare the five submitted workbooks against newly generated workbooks."""

from argparse import ArgumentParser
from datetime import datetime
from itertools import zip_longest
from pathlib import Path

from openpyxl import load_workbook

NAMES = ("result1", "result2", "result3", "result4-2", "result4-3")


def normalized(value):
    return value.date() if isinstance(value, datetime) else value


def compare_book(expected_path: Path, actual_path: Path, tolerance: float):
    expected = load_workbook(expected_path, read_only=True, data_only=True)
    actual = load_workbook(actual_path, read_only=True, data_only=True)
    try:
        if expected.sheetnames != actual.sheetnames:
            raise ValueError("sheet names differ")
        checked = mismatches = 0
        worst = 0.0
        for sheet_name in expected.sheetnames:
            left = expected[sheet_name].iter_rows(values_only=True)
            right = actual[sheet_name].iter_rows(values_only=True)
            for left_row, right_row in zip_longest(left, right, fillvalue=()):
                for a, b in zip_longest(left_row, right_row, fillvalue=None):
                    checked += 1
                    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                        difference = abs(a - b)
                        worst = max(worst, difference)
                        same = difference <= tolerance
                    else:
                        same = normalized(a) == normalized(b)
                    if not same:
                        mismatches += 1
        return checked, mismatches, worst
    finally:
        expected.close()
        actual.close()


def main():
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--expected", required=True, type=Path)
    parser.add_argument("--actual", required=True, type=Path)
    parser.add_argument("--tolerance", type=float, default=1e-5)
    args = parser.parse_args()
    if args.tolerance < 0:
        parser.error("--tolerance must be non-negative")
    failed = False
    for name in NAMES:
        try:
            checked, mismatches, worst = compare_book(
                args.expected / f"{name}.xlsx",
                args.actual / f"{name}.xlsx",
                args.tolerance,
            )
        except (OSError, ValueError) as error:
            print(f"{name}: FAIL ({error})")
            failed = True
            continue
        print(f"{name}: cells={checked}, mismatches={mismatches}, max_numeric_diff={worst:.6g}")
        failed |= mismatches > 0
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
