# -*- coding: utf-8 -*-
from pathlib import Path
import json,datetime
from openpyxl import load_workbook
from solve import ROOT

def interval_label(index):
    return f'{index // 6}:{index % 6 * 10:02d}-{(index + 1) // 6}:{(index + 1) % 6 * 10:02d}'

def main():
    import argparse
    ap=argparse.ArgumentParser();ap.add_argument('--xlsx-dir',type=Path,default=ROOT/'results');args=ap.parse_args()
    payload=json.loads((ROOT/'results/export_values.json').read_text(encoding="utf-8"));checks={}
    for name,sheets in payload.items():
        w=load_workbook(args.xlsx_dir/f'{name}.xlsx',data_only=True)
        template=load_workbook(ROOT/'data/templates'/f'{name}.xlsx',data_only=True)
        count=0;worst=0.
        assert set(w.sheetnames)==set(sheets)
        for sn,rows in sheets.items():
            s=w[sn]
            for i,row in enumerate(rows,2):
                for j,v in enumerate(row,1):
                    actual=s.cell(i,j).value
                    if name=='result1' and sn=='计划购电量' and j==1:
                        assert actual==interval_label(i-2),(name,sn,i,j,actual)
                        count+=1
                        continue
                    if isinstance(v,(int,float)):
                        assert isinstance(actual,(int,float)),(name,sn,i,j,actual)
                        diff=abs(actual-v);worst=max(worst,diff)
                        assert diff<1e-5,(name,sn,i,j,diff)
                    elif v is None:assert actual is None,(name,sn,i,j,actual)
                    elif isinstance(v,str) and v.startswith('2025-'):
                        assert actual.date()==datetime.date.fromisoformat(v),(name,sn,i,j,actual)
                    else:assert actual==v,(name,sn,i,j,actual,v)
                    count+=1
            if sn in ['计划购电量','调整购电量']:
                if name=='result1':
                    assert s['A2'].value==interval_label(0) and s['A145'].value==interval_label(143)
                else:
                    for col in range(2,146):
                        assert s.cell(1,col).value==interval_label(col-2),(name,sn,col)
        checks[name]={'checked_cells':count,'max_numeric_difference':worst,'status':'PASS'}
    (ROOT/'results/export_validation.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2), encoding="utf-8")
    print(json.dumps(checks,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
